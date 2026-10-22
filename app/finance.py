"""COD money: reconcile the carrier's payout file with delivered orders, and the monthly statement.

The carrier collects the cash and pays the merchant later, minus its fees, with a file listing the
parcels. Mrigel matches each line to an order (tracking, then order reference, then a unique phone),
marks it paid and delivered, and lists what does not add up: unknown parcels, amounts that differ
from the order total, parcels paid twice, and delivered orders that were never paid.
"""

import csv
import io
import re
import unicodedata
from datetime import datetime, timedelta

from sqlalchemy import select

from .algeria import normalize_phone
from .models import Order, Payout, utcnow
from .services import as_utc, set_stage

COLUMNS = {
    "tracking": ["tracking", "tracking number", "code", "code colis", "code de suivi", "n suivi", "numero de suivi",
                 "suivi", "colis", "رقم التتبع", "التتبع"],
    "reference": ["reference", "ref", "order", "order id", "commande", "n commande", "id commande", "رقم الطلبية", "المرجع"],
    "phone": ["phone", "telephone", "tel", "mobile", "الهاتف", "رقم الهاتف"],
    "amount": ["amount", "montant", "montant cod", "montant encaisse", "encaisse", "cod", "prix", "total",
               "المبلغ", "المبلغ المحصل"],
    "fee": ["fee", "fees", "frais", "frais de livraison", "frais livraison", "tarif", "cout", "commission", "الرسوم",
            "سعر التوصيل"],
}


def _fold(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    text = re.sub(r"[_\-°º.#:]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


_FOLDED = {field: {_fold(n) for n in names} for field, names in COLUMNS.items()}


def _money(value) -> int | None:
    text = re.sub(r"[^\d,.\-]", "", str(value or ""))
    if not text:
        return None
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    text = text.replace(",", "")
    try:
        return round(float(text))
    except ValueError:
        return None


def read_rows(data: bytes, filename: str = "") -> list[dict]:
    """CSV (any of , ; tab, UTF-8 or Windows-1256) or Excel .xlsx -> [{tracking, reference, phone, amount, fee}]."""
    if filename.lower().endswith(".xlsx") or data[:2] == b"PK":
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        table = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    else:
        for encoding in ("utf-8-sig", "cp1256", "latin-1"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        # carriers' files use ; or tab with decimal commas, so count separators instead of sniffing
        head = text.splitlines()[:8]
        delimiter = max((";", "\t", ","), key=lambda d: sum(line.count(d) for line in head))
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    # the header is the first row that names at least two known columns (files often start with a title)
    for start, header in enumerate(table[:10]):
        mapping = {}
        for index, cell in enumerate(header):
            for field, names in _FOLDED.items():
                if field not in mapping.values() and _fold(cell) in names:
                    mapping[index] = field
        if len(mapping) >= 2 and "amount" in mapping.values():
            break
    else:
        return []
    rows = []
    for raw in table[start + 1:]:
        record = {field: (raw[i].strip() if i < len(raw) else "") for i, field in mapping.items()}
        if not any(record.values()):
            continue
        record["amount"] = _money(record.get("amount"))
        record["fee"] = _money(record.get("fee")) or 0
        rows.append(record)
    return rows


def _find(db, store_id: str, row: dict) -> Order | None:
    tracking = (row.get("tracking") or "").strip()
    if tracking:
        order = db.scalar(select(Order).where(Order.store_id == store_id, Order.tracking == tracking))
        if order:
            return order
    reference = (row.get("reference") or "").strip()
    if reference:
        order = db.scalar(select(Order).where(Order.store_id == store_id, Order.external_id == reference))
        if order:
            return order
        if re.fullmatch(r"[0-9A-Fa-f]{8}", reference):  # the short reference our carrier export writes
            candidates = db.scalars(select(Order).where(Order.store_id == store_id,
                                                        Order.id.like(reference.lower() + "%"))).all()
            if len(candidates) == 1:
                return candidates[0]
    if row.get("phone"):
        phone, valid = normalize_phone(row["phone"])
        if valid:
            candidates = db.scalars(select(Order).where(Order.store_id == store_id, Order.phone == phone,
                                                        Order.stage.in_(("shipped", "delivered")),
                                                        Order.paid_at.is_(None))).all()
            if len(candidates) == 1:
                return candidates[0]
    return None


def reconcile(db, store, data: bytes, filename: str = "", carrier: str = "", user_id: str | None = None) -> Payout:
    rows = read_rows(data, filename)
    payout = Payout(store_id=store.id, carrier=carrier, filename=filename[:200], rows=len(rows), created_by_id=user_id)
    db.add(payout)
    db.flush()
    issues = []
    now = utcnow()
    for number, row in enumerate(rows, start=1):
        label = row.get("tracking") or row.get("reference") or row.get("phone") or f"#{number}"
        amount = row.get("amount")
        order = _find(db, store.id, row)
        if order is None:
            issues.append({"kind": "not_found", "line": number, "label": label, "amount": amount or 0})
            continue
        if order.paid_at is not None and order.payout_id != payout.id:
            issues.append({"kind": "already_paid", "line": number, "label": label, "order_id": order.id,
                           "amount": amount or 0})
            continue
        if order.stage == "shipped":
            set_stage(db, order, "delivered", user_id=user_id, reason="payout")
        elif order.stage != "delivered":
            issues.append({"kind": "not_delivered", "line": number, "label": label, "order_id": order.id,
                           "stage": order.stage, "amount": amount or 0})
            continue
        order.paid_at = now
        order.payout_id = payout.id
        order.cod_collected_dzd = amount if amount is not None else order.total_dzd
        order.carrier_fee_dzd = row.get("fee") or 0
        if carrier and not order.carrier:
            order.carrier = carrier
        payout.matched += 1
        payout.collected_dzd += order.cod_collected_dzd or 0
        payout.fees_dzd += order.carrier_fee_dzd or 0
        if amount is not None and amount != order.total_dzd:
            issues.append({"kind": "amount_mismatch", "line": number, "label": label, "order_id": order.id,
                           "amount": amount, "expected": order.total_dzd})
    payout.issues = issues
    return payout


def summary(db, store, days: int = 30, unpaid_after_days: int = 7) -> dict:
    now = utcnow()
    orders = db.scalars(select(Order).where(Order.store_id == store.id,
                                            Order.stage.in_(("shipped", "delivered", "returned")))).all()
    delivered_unpaid = [o for o in orders if o.stage == "delivered" and o.paid_at is None]
    late = [o for o in delivered_unpaid if as_utc(o.closed_at or o.updated_at) <= now - timedelta(days=unpaid_after_days)]
    since = now - timedelta(days=days)
    paid = [o for o in orders if o.paid_at is not None and as_utc(o.paid_at) >= since]
    in_transit = [o for o in orders if o.stage == "shipped"]
    collected = sum(o.cod_collected_dzd or 0 for o in paid)
    fees = sum(o.carrier_fee_dzd or 0 for o in paid)
    return {
        "days": days,
        "expected_dzd": sum(o.total_dzd for o in delivered_unpaid),
        "expected_n": len(delivered_unpaid),
        "in_transit_dzd": sum(o.total_dzd for o in in_transit),
        "in_transit_n": len(in_transit),
        "collected_dzd": collected,
        "fees_dzd": fees,
        "net_dzd": collected - fees,
        "paid_n": len(paid),
        "late": sorted(late, key=lambda o: as_utc(o.closed_at or o.updated_at))[:50],
        "late_n": len(late),
        "unpaid_after_days": unpaid_after_days,
    }


def month_bounds(month: str) -> tuple[datetime, datetime]:
    """'2026-10' -> start and end in Algiers time, as aware datetimes."""
    from zoneinfo import ZoneInfo

    tz = ZoneInfo("Africa/Algiers")
    year, mon = (int(x) for x in month.split("-"))
    start = datetime(year, mon, 1, tzinfo=tz)
    end = datetime(year + (mon == 12), 1 if mon == 12 else mon + 1, 1, tzinfo=tz)
    return start, end


def month_figures(db, store, month: str) -> dict:
    from datetime import timezone

    local_start, local_end = month_bounds(month)
    start, end = local_start.astimezone(timezone.utc), local_end.astimezone(timezone.utc)
    created = db.scalars(select(Order).where(Order.store_id == store.id, Order.created_at >= start,
                                             Order.created_at < end)).all()
    closed = db.scalars(select(Order).where(Order.store_id == store.id, Order.closed_at >= start,
                                            Order.closed_at < end)).all()
    paid = db.scalars(select(Order).where(Order.store_id == store.id, Order.paid_at >= start,
                                          Order.paid_at < end)).all()
    delivered = [o for o in closed if o.stage == "delivered"]
    returned = [o for o in closed if o.stage == "returned"]
    from .store_settings import get_settings

    return_cost = int(get_settings(store).get("return_cost_dzd") or 0)
    by_day: dict[str, dict] = {}
    for order in delivered:
        day = as_utc(order.closed_at).astimezone(local_start.tzinfo).strftime("%d/%m")
        row = by_day.setdefault(day, {"n": 0, "dzd": 0})
        row["n"] += 1
        row["dzd"] += order.total_dzd
    return {
        "month": month, "orders": len(created),
        "confirmed": sum(1 for o in created if o.confirmed_at is not None),
        "delivered_n": len(delivered), "delivered_dzd": sum(o.total_dzd for o in delivered),
        "returned_n": len(returned), "returns_cost_dzd": len(returned) * return_cost,
        "collected_dzd": sum(o.cod_collected_dzd or 0 for o in paid),
        "fees_dzd": sum(o.carrier_fee_dzd or 0 for o in paid),
        "paid_n": len(paid),
        "by_day": sorted(by_day.items()),
    }
